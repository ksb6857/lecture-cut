# -*- coding: utf-8 -*-
"""pptx 에 '클릭할 때 · 나타내기(밝기 변화) · 0.5초' 애니메이션을 단계별로 써 넣는다.

한 클릭 = 대본 한 문단. 자동 재생은 넣지 않는다. 녹화할 때 발표자가 넘기는 대로만
나타나야 하기 때문이다. 효과는 한 가지만 쓴다(납품 가이드 예: 1~2가지, 0.5초, 클릭할 때,
장당 3~6단계). presetID 10 / presetClass entr / filter fade.

다른 교안 작업에서 쓰고 파워포인트에서 열어 확인한 코드를 옮겨 왔다. 바꾼 것은 기본 길이 0.5초와 계획 파일로 한꺼번에 넣는 run() 뿐이다.

**원본을 고치지 않는다.** 사본에 써서 따로 저장한다. 넣은 뒤 반드시
발표 자료 비교 도구(글자·글꼴·위치·크기를 장별로 잰다)로 애니메이션 말고 바뀐 것이 0인지, 파워포인트로 한 장이라도
렌더되는지 본다(python-pptx 로 저장한 파일이 파워포인트에서 안 열린 적이 있다).

계획 파일:
  {"slides": {"6": [[12, 14, 17, 20, 23], [26, 28, 31, 34, 37], [40]], ...}}
  바깥 리스트 하나가 클릭 한 번. 안의 숫자는 도형 id (최상위 도형).

사용: python ppt_anim.py <원본.pptx> <계획.json> <새이름.pptx>
"""
import json
import sys
from pathlib import Path

from lxml import etree

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

_EFFECT = """
<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0"
              fill="hold" grpId="0" nodeType="{node}">
  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
  <p:childTnLst>
    <p:set>
      <p:cBhvr>
        <p:cTn id="{sid1}" dur="1" fill="hold">
          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
        </p:cTn>
        <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
        <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
      </p:cBhvr>
      <p:to><p:strVal val="visible"/></p:to>
    </p:set>
    <p:animEffect transition="in" filter="fade">
      <p:cBhvr>
        <p:cTn id="{sid2}" dur="{dur}"/>
        <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
      </p:cBhvr>
    </p:animEffect>
  </p:childTnLst>
</p:cTn></p:par>"""

_CLICK = """
<p:par><p:cTn id="{cid}" fill="hold">
  <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
  <p:childTnLst>
    <p:par><p:cTn id="{cid2}" fill="hold">
      <p:stCondLst><p:cond delay="0"/></p:stCondLst>
      <p:childTnLst>{effects}</p:childTnLst>
    </p:cTn></p:par>
  </p:childTnLst>
</p:cTn></p:par>"""

_TIMING = """<p:timing xmlns:p="{p}" xmlns:a="{a}">
<p:tnLst>
 <p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
  <p:childTnLst>
   <p:seq concurrent="1" nextAc="seek">
    <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
     <p:childTnLst>{clicks}</p:childTnLst>
    </p:cTn>
    <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
    <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
   </p:seq>
  </p:childTnLst>
 </p:cTn></p:par>
</p:tnLst>
</p:timing>"""


def clicks(slide, groups, dur=500):
    """groups: [[도형, 도형, ...], ...] — 바깥 리스트 하나가 클릭 한 번.
    한 그룹 안의 첫 도형이 클릭 효과, 나머지는 같이 나타난다."""
    groups = [[s for s in g if s is not None] for g in groups]
    groups = [g for g in groups if g]
    if not groups:
        return slide
    nid = [3]

    def take():
        nid[0] += 1
        return nid[0] - 1

    click_xml = []
    for g in groups:
        eff = []
        for k, sh in enumerate(g):
            eff.append(_EFFECT.format(
                cid=take(), sid1=take(), sid2=take(), spid=sh.shape_id, dur=dur,
                node="clickEffect" if k == 0 else "withEffect"))
        click_xml.append(_CLICK.format(cid=take(), cid2=take(), effects="".join(eff)))
    timing = etree.fromstring(_TIMING.format(p=P, a=A, clicks="".join(click_xml)))
    sld = slide._element
    old = sld.find("{%s}timing" % P)
    if old is not None:
        sld.remove(old)
    # p:timing 은 p:clrMapOvr 뒤, p:extLst 앞에 와야 한다(스키마 순서)
    ext = sld.find("{%s}extLst" % P)
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)
    return slide


def count(slide):
    t = slide._element.find("{%s}timing" % P)
    return 0 if t is None else len(t.findall(".//{%s}cTn[@nodeType='clickEffect']" % P))


def run(src, plan, out, dur=500):
    from pptx import Presentation
    prs = Presentation(src)
    done = {}
    for key, steps in plan["slides"].items():
        s = prs.slides[int(key) - 1]
        by_id = {sh.shape_id: sh for sh in s.shapes}
        miss = [i for g in steps for i in g if i not in by_id]
        if miss:
            raise SystemExit(f"S{key}: 없는 도형 id {miss}")
        clicks(s, [[by_id[i] for i in g] for g in steps], dur)
        done[int(key)] = count(s)
    prs.save(out)
    return done


if __name__ == "__main__":
    plan = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    d = run(sys.argv[1], plan, sys.argv[3])
    print(f"{len(d)}장에 애니메이션 · 클릭 {sum(d.values())}번 → {sys.argv[3]}")
