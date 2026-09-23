# -*- coding: utf-8 -*-
"""슬라이드 도형을 '한 번에 나타날 묶음'으로 가른다 — 나타내기 애니메이션의 단계 후보.

강의 화면이 한 장에 오래 멈춰 있으면 지금 어디를 말하는지 알기 어렵다.
설명하는 순서대로 항목이 하나씩 나타나야 한다(납품 가이드 예: 한 단계 = 대본 한 문단,
장당 3~6단계, 나타내기 한 종류, 0.5초).

가르는 법
  - **바탕**: 배경·액자 틀·제목·작은 장식은 처음부터 떠 있다
  - **카드**: 다른 도형을 품는 큰 도형이 카드다. 중심이 카드 안(여유 12px)에 든
    도형은 그 카드에 붙는다(카드 위에 걸친 번호 동그라미 포함)
  - **순서**: 줄(세로로 겹치는 것) 단위로 위→아래, 줄 안에서 왼쪽→오른쪽

자동으로 가른 뒤 사람이 내레이션을 보고 합치거나 쪼갠다. 결과는 계획 파일로 넘긴다.

사용: python slide_blocks.py <pptx> [장번호...]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def text_of(sh):
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        return " ".join(t for t in (text_of(s) for s in sh.shapes) if t)
    if sh.has_text_frame:
        return sh.text_frame.text.replace("\n", " ").strip()
    return ""


def geom(prs, sh):
    W, H = prs.slide_width, prs.slide_height
    x0, y0 = sh.left / W * 1920, sh.top / H * 1080
    return [round(x0), round(y0), round(x0 + sh.width / W * 1920), round(y0 + sh.height / H * 1080)]


def is_base(i, sh, g, text):
    x0, y0, x1, y1 = g
    w, h = x1 - x0, y1 - y0
    if i < 2:
        return True                                  # 배경, 액자 틀
    if w >= 1700 and h >= 900:
        return True
    if text and y1 <= 300 and h <= 140:              # 제목
        return True
    if y1 <= 300 and not text:                       # 제목 옆 장식·꼬리표
        return True
    if not text and w * h < 20000 and sh.shape_type in (MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE):
        return True                                  # 작은 장식
    return False


def blocks(prs, idx):
    s = prs.slides[idx - 1]
    items = []
    for i, sh in enumerate(s.shapes):
        g = geom(prs, sh)
        t = text_of(sh)
        items.append({"id": sh.shape_id, "name": sh.name, "g": g, "text": t,
                      "base": is_base(i, sh, g, t),
                      "pic": sh.shape_type == MSO_SHAPE_TYPE.PICTURE})
    body = [it for it in items if not it["base"]]

    def area(g):
        return (g[2] - g[0]) * (g[3] - g[1])

    def inside(a, b, tol=12):                         # a 의 중심이 b 안에?
        cx, cy = (a["g"][0] + a["g"][2]) / 2, (a["g"][1] + a["g"][3]) / 2
        return (b["g"][0] - tol <= cx <= b["g"][2] + tol and b["g"][1] - tol <= cy <= b["g"][3] + tol
                and area(b["g"]) > area(a["g"]) * 1.5)

    # 가장 큰 것부터 카드로 삼고, 안에 든 것을 붙인다
    body.sort(key=lambda it: -area(it["g"]))
    groups = []
    for it in body:
        host = next((g for g in groups if inside(it, g[0])), None)
        if host:
            host.append(it)
        else:
            groups.append([it])
    # 캡션(사진 바로 아래 작은 글)은 사진에 붙인다
    for g in list(groups):
        if len(g) == 1 and g[0]["text"].startswith("이미지"):
            pic = next((h for h in groups if h is not g and h[0]["pic"]
                        and abs(h[0]["g"][3] - g[0]["g"][1]) < 30), None)
            if pic:
                pic.extend(g)
                groups.remove(g)

    def box(g):
        return [min(i["g"][0] for i in g), min(i["g"][1] for i in g),
                max(i["g"][2] for i in g), max(i["g"][3] for i in g)]

    # 줄 단위 정렬: 세로 범위가 절반 넘게 겹치면 같은 줄
    gs = [{"ids": [i["id"] for i in g], "box": box(g),
           "text": " / ".join(i["text"] for i in sorted(g, key=lambda i: (i["g"][1], i["g"][0])) if i["text"])}
          for g in groups]
    gs.sort(key=lambda b: (b["box"][1], b["box"][0]))
    rows = []
    for b in gs:
        y0, y1 = b["box"][1], b["box"][3]
        row = next((r for r in rows if min(y1, r["y1"]) - max(y0, r["y0"]) > 0.5 * min(y1 - y0, r["y1"] - r["y0"])), None)
        if row:
            row["items"].append(b)
            row["y0"], row["y1"] = min(row["y0"], y0), max(row["y1"], y1)
        else:
            rows.append({"y0": y0, "y1": y1, "items": [b]})
    ordered = [b for r in sorted(rows, key=lambda r: r["y0"])
               for b in sorted(r["items"], key=lambda b: b["box"][0])]
    base = [it["id"] for it in items if it["base"]]
    title = next((it["text"] for it in items if it["base"] and it["text"]), "")
    return {"slide": idx, "title": title, "base": base, "steps": ordered}


def main(argv):
    prs = Presentation(argv[0])
    only = [int(x) for x in argv[1:]] or range(1, len(prs.slides) + 1)
    for i in only:
        b = blocks(prs, i)
        print(f"\nS{i:02d} {b['title'][:40]}  (바탕 {len(b['base'])})")
        for k, st in enumerate(b["steps"], 1):
            print(f"   {k}. {st['box']} ids{st['ids']} | {st['text'][:70]}")


if __name__ == "__main__":
    main(sys.argv[1:])
